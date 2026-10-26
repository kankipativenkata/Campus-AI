import io
import wave

import pytest

from core import chat_memory, config, db, interview as I, placement as P, projects as PJ, rag, roadmap as R, timetable as T


# ---------------- conversation memory ----------------
def test_followup_detection():
    assert rag.is_followup("And can it be condoned?")
    assert not rag.is_followup("How many casual leaves do faculty get per year?")


def test_followup_is_resolved_offline():
    hist = [{"q": "What is the minimum attendance for exam eligibility?", "a": "75%"}]
    res = rag.answer("And can it be condoned?", "student", history=hist)
    assert res["found"] and "condonation" in res["answer"].lower()
    assert "attendance" in res["standalone"].lower()


def test_followup_without_history_is_not_guessed():
    assert rag.answer("and can it be condoned?", "student")["standalone"] == "and can it be condoned?"


def test_llm_condense_used_when_available(monkeypatch):
    from core import llm
    monkeypatch.setattr(llm, "available", lambda: True)
    monkeypatch.setattr(llm, "generate", lambda *a, **k: "Can attendance shortage be condoned?")
    monkeypatch.setattr(llm, "embed", lambda *a, **k: None)
    out = rag.condense_question("and can it?", [{"q": "attendance rule?", "a": "75%"}])
    assert out == "Can attendance shortage be condoned?"


def test_memory_persists_and_resets():
    u = "memuser"
    for i in range(3):
        db.insert(db.chat_messages, username=u, question=f"q{i}", answer=f"a{i}", sources="[]", found=True, feedback=None, created_at=db.now())
    assert [h["q"] for h in chat_memory.load_history(u)] == ["q0", "q1", "q2"]
    assert chat_memory.turns_for_prompt(chat_memory.load_history(u), 2) == [{"q": "q1", "a": "a1"}, {"q": "q2", "a": "a2"}]
    chat_memory.reset(u)
    assert chat_memory.load_history(u) == []
    db.insert(db.chat_messages, username=u, question="after", answer="x", sources="[]", found=True, feedback=None, created_at=db.now())
    assert [h["q"] for h in chat_memory.load_history(u)] == ["after"]
    assert chat_memory.load_history("someone_else") == []


# ---------------- timetable ----------------
def _cfg(**kw):
    base = {"days": T.DAYS_ALL[:5], "periods": T.make_periods("09:00", 8, 50, {4: 60}), "theory_rooms": ["R1", "R2"],
            "lab_rooms": ["L1", "L2"], "unavailable": T.SAMPLE_UNAVAILABLE}
    base.update(kw)
    return base


def test_periods_have_break():
    p = T.make_periods("09:00", 8, 50, {4: 60})
    assert [x["label"] for x in p].count("Break") == 1 and p[4]["label"] == "Break" and p[4]["start"] == "12:20"


def test_sample_timetable_is_complete_and_conflict_free():
    tt = T.solve(_cfg(), T.SAMPLE_COURSES, seed=1)
    need = sum(c["theory_hours"] + c["lab_hours"] for c in T.SAMPLE_COURSES)
    assert not tt["unplaced"] and tt["conflicts"] == []
    assert sum(p["len"] for p in tt["placements"]) == need


def test_labs_are_blocks_in_lab_rooms_not_across_break():
    tt = T.solve(_cfg(), T.SAMPLE_COURSES, seed=2)
    labs = [p for p in tt["placements"] if p["kind"] == "lab"]
    assert labs and all(p["len"] == 2 and p["room"] in ("L1", "L2") for p in labs)
    for p in labs:
        assert all(tt["periods"][i]["teach"] for i in range(p["start"], p["start"] + p["len"]))


def test_lab_assistant_is_never_double_booked():
    tt = T.solve(_cfg(), T.SAMPLE_COURSES, seed=3)
    assert tt["conflicts"] == []
    slots = [(p["day"], i) for p in tt["placements"] if p.get("assistant") == "Mr. Kiran" for i in range(p["start"], p["start"] + p["len"])]
    assert len(slots) == len(set(slots))


def test_unavailability_respected():
    tt = T.solve(_cfg(), T.SAMPLE_COURSES, seed=4)
    for p in tt["placements"]:
        if p["faculty"] == "Dr. Rao" and p["day"] == "Mon":
            assert p["start"] >= 2


def test_validator_detects_a_forced_clash():
    tt = T.solve(_cfg(), T.SAMPLE_COURSES, seed=1)
    bad = dict(tt, placements=tt["placements"] + [dict(tt["placements"][0])])
    assert T.validate(bad)


def test_overloaded_faculty_is_reported_not_hidden():
    courses = [{"course": f"C{i}", "name": "x", "faculty": "Dr. One", "section": f"S{i}", "theory_hours": 10, "lab_hours": 0, "lab_block": 2, "lab_assistant": ""} for i in range(6)]
    tt = T.solve(_cfg(unavailable=[]), courses, seed=0, restarts=5)
    assert tt["unplaced"] and "No free slot" in tt["unplaced"][0]["reason"] and tt["conflicts"] == []


def test_no_lab_rooms_gives_clear_reason():
    tt = T.solve(_cfg(lab_rooms=[]), T.SAMPLE_COURSES[:1], seed=0, restarts=2)
    assert any("No lab rooms" in u["reason"] for u in tt["unplaced"])


def test_same_seed_is_deterministic():
    a = T.solve(_cfg(), T.SAMPLE_COURSES, seed=7, restarts=6)["placements"]
    b = T.solve(_cfg(), T.SAMPLE_COURSES, seed=7, restarts=6)["placements"]
    assert a == b


def test_html_export_escapes_user_text():
    courses = [{"course": "<script>x</script>", "name": "n", "faculty": "<b>F</b>", "section": "A", "theory_hours": 1, "lab_hours": 0, "lab_block": 2, "lab_assistant": ""}]
    tt = T.solve(_cfg(unavailable=[]), courses, seed=0, restarts=2)
    html = T.to_html(tt, "<img src=x>")
    assert "<script>x" not in html and "&lt;script&gt;" in html and "<img src=x>" not in html


def test_share_code_roundtrip_and_owner_only_delete():
    tt = T.solve(_cfg(), T.SAMPLE_COURSES, seed=1, restarts=3)
    code = T.save("fac_owner", "My TT", tt)
    assert T.load(code)["title"] == "My TT" and T.load("nope") is None
    T.unshare(code, "someone_else")
    assert T.load(code) is not None
    T.unshare(code, "fac_owner")
    assert T.load(code) is None


def test_excel_export_has_sheets():
    import openpyxl
    tt = T.solve(_cfg(), T.SAMPLE_COURSES, seed=1, restarts=3)
    wb = openpyxl.load_workbook(io.BytesIO(T.to_excel(tt)))
    assert "All sessions" in wb.sheetnames and any(n.startswith("Sec ") for n in wb.sheetnames)


# ---------------- roadmap ----------------
CAT, GOALS = R.load_catalog(), R.load_goals()


@pytest.mark.parametrize("goal", list(GOALS))
def test_roadmap_respects_prereqs_and_credit_cap(goal):
    p = R.plan(GOALS[goal], CAT, set(), 16, 10)
    done = set()
    for sem in p["semesters"]:
        assert sem["credits"] <= 16
        for c in sem["courses"]:
            assert all(x in done for x in CAT[c]["prereqs"]), (goal, c)
        done |= set(sem["courses"])
    assert not p["unscheduled"] and p["covered_pct"] > 70


def test_completed_courses_reduce_plan():
    full = R.plan(GOALS["Data Analyst"], CAT, set())
    part = R.plan(GOALS["Data Analyst"], CAT, {"CSE1001", "CSE2003"})
    assert part["total_credits"] < full["total_credits"]
    assert "CSE1001" not in {c for s in part["semesters"] for c in s["courses"]}


def test_uncovered_skills_are_reported():
    assert "rag" in R.plan(GOALS["AI / LLM Engineer"], CAT, set())["uncovered"]


def test_roadmap_exports():
    p = R.plan(GOALS["ML Engineer"], CAT, set())
    assert R.to_dot(p, CAT, set()).startswith("digraph") and "Semester 1" in R.to_markdown("ML Engineer", p, CAT)


# ---------------- placement ----------------
DATA = P.load()
CO = DATA["companies"][0]
S = {"cgpa": 7.8, "backlogs": 0, "projects": 1, "internships": 0, "problems": 60, "comm": 3, "skills": ["python"]}


def test_ineligible_students_get_zero_with_reasons():
    e = P.estimate({**S, "cgpa": 6.0}, CO, DATA["prior"])
    assert not e["eligible"] and e["mean"] == 0 and "CGPA" in e["reasons"][0]
    assert not P.estimate({**S, "backlogs": 2}, CO, DATA["prior"])["eligible"]


def test_probability_increases_with_skills_and_is_a_valid_interval():
    lo = P.estimate(S, CO, DATA["prior"])
    hi = P.estimate({**S, "skills": list(CO["required_skills"]), "projects": 3, "internships": 1, "problems": 400}, CO, DATA["prior"])
    assert hi["mean"] > lo["mean"] * 2
    for e in (lo, hi):
        assert 0 <= e["low"] <= e["mean"] <= e["high"] <= 1


def test_less_history_means_wider_interval():
    big, small = DATA["companies"][1], DATA["companies"][4]
    a, b = P.estimate(S, big, DATA["prior"]), P.estimate(S, small, DATA["prior"])
    assert (b["high"] - b["low"]) / max(b["mean"], 1e-9) > (a["high"] - a["low"]) / max(a["mean"], 1e-9)


def test_backlog_lowers_probability():
    c = DATA["companies"][1]
    assert P.estimate({**S, "backlogs": 1}, c, DATA["prior"])["mean"] < P.estimate(S, c, DATA["prior"])["mean"]


def test_improvements_sorted_and_actionable():
    imp = P.improvements(S, CO, DATA["prior"])
    assert imp and [x["gain_pp"] for x in imp] == sorted((x["gain_pp"] for x in imp), reverse=True)
    assert any(x["action"].startswith("Learn dsa") for x in imp)
    assert P.improvements({**S, "cgpa": 5}, CO, DATA["prior"]) == []


# ---------------- interview ----------------
BANK = I.load_bank()
Q = BANK["Viva: Data Structures"]["questions"][0]
GOOD = "A stack is a linear data structure that follows LIFO, last in first out. We push and pop elements at the top. It is used for undo, recursion and expression evaluation."
PARTIAL = "A stack lets you push items on the top and pop them off again when you need them back for something later in the program."
POOR = "Um, like, basically it is a thing, you know, um like a queue I think."


def test_good_beats_partial_beats_poor():
    g, p, r = (I.score_answer(Q, t) for t in (GOOD, PARTIAL, POOR))
    assert g["overall"] > p["overall"] > r["overall"]
    assert g["correctness"] >= 9 and r["correctness"] <= 2


def test_filler_and_clarity_signals():
    assert I.filler_count("Um, so like, you know, it works") >= 3
    assert I.clarity_heuristic(POOR)["score"] < I.clarity_heuristic(GOOD)["score"]
    assert I.clarity_heuristic("")["score"] == 0
    fast = I.clarity_heuristic(GOOD * 3, duration_s=20)
    assert fast["wpm"] > 190 and any("fast" in t.lower() for t in fast["tips"])


def test_wav_duration():
    buf = io.BytesIO()
    with wave.open(buf, "wb") as w:
        w.setnchannels(1); w.setsampwidth(2); w.setframerate(16000); w.writeframes(b"\0\0" * 32000)
    assert I.wav_duration(buf.getvalue()) == pytest.approx(2.0)
    assert I.wav_duration(b"not audio") is None


def test_pick_questions_and_bank_integrity():
    assert len(I.pick_questions(BANK, "Viva: Databases", 2, seed=1)) == 2
    for cat in BANK.values():
        for q in cat["questions"]:
            assert q["key_points"] and all(k["keywords"] for k in q["key_points"])


def test_ai_scoring_blends_with_rules(monkeypatch):
    from core import llm
    monkeypatch.setattr(llm, "available", lambda: True)
    monkeypatch.setattr(llm, "generate_json", lambda *a, **k: {"correctness": 10, "clarity": 8, "feedback": "Good.", "model_answer": "x", "mistakes": []})
    r = I.score_answer(Q, GOOD)
    assert r["source"] == "AI + rules" and r["feedback"] == "Good." and 0 <= r["overall"] <= 10


def test_transcribe_without_key_fails_gracefully():
    from core import llm
    with pytest.raises(llm.LLMError):
        llm.transcribe(b"x")


# ---------------- project showcase ----------------
def test_rule_tags_use_word_boundaries():
    t = PJ.rule_tags("Built with Python and Flask, uses a digit classifier and github")
    assert {"python", "flask", "git"} <= set(t["tech"]) and "java" not in t["tech"]


def test_auto_tag_offline():
    t = PJ.auto_tag("Hostel monitor", "ESP32 and Arduino IoT sensors with a dashboard", use_ai=False)
    assert "esp32" in t["tech"] and "IoT / Embedded" in t["domains"] and t["source"] == "rules"


def test_similar_projects_rank_related_first():
    allp = PJ.load_projects()
    target = next(p for p in allp if p["title"].startswith("Brain tumor"))
    sims = PJ.similar_projects(target, [p for p in allp if p["id"] != target["id"]], 3)
    assert sims[0]["title"].startswith("Chest X-ray") and sims[0]["score"] > sims[-1]["score"]


def test_teammates_are_complementary_and_consent_based():
    allp, profs = PJ.load_projects(), PJ.load_profiles()
    target = next(p for p in allp if p["title"].startswith("Syllabus Q&A"))
    target = {**target, "my_skills": ["python", "llm", "fastapi"]}
    sims = PJ.similar_projects(target, [p for p in allp if p["id"] != target["id"]], 5)
    names = [t["username"] for t in PJ.teammate_suggestions(target, sims, profs, "seed_e")]
    assert "seed_d" not in names and "seed_e" not in names and names


def test_publish_and_delete_permissions():
    tags = PJ.auto_tag("Quiz app", "A Flask quiz app for students", use_ai=False)
    pid = PJ.add_project("owner1", "Owner", "Quiz app", "A Flask quiz app for students", "", tags, embed=False)
    PJ.delete_project(pid, "intruder")
    assert any(p["id"] == pid for p in PJ.load_projects())
    PJ.delete_project(pid, "owner1")
    assert not any(p["id"] == pid for p in PJ.load_projects())
    pid = PJ.add_project("owner1", "Owner", "Quiz app", "x", "", tags, embed=False)
    PJ.delete_project(pid, "admin", is_admin=True)
    assert not any(p["id"] == pid for p in PJ.load_projects())
