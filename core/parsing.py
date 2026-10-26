"""Turn uploaded files into [(page, text)]. PDF, DOCX, XLSX/CSV, TXT/MD, and images (via vision OCR)."""
from __future__ import annotations

import io
from pathlib import Path

from core import llm

IMAGE_EXT = {".png", ".jpg", ".jpeg", ".webp"}
TEXT_EXT = {".txt", ".md"}
SUPPORTED = sorted(IMAGE_EXT | TEXT_EXT | {".pdf", ".docx", ".csv", ".xlsx"})


def _df_to_text(df) -> str:
    lines = []
    for _, row in df.iterrows():
        cells = [f"{c}: {v}" for c, v in row.items() if str(v).strip() and str(v) != "nan"]
        if cells:
            lines.append("; ".join(cells))
    return "\n".join(lines)


def extract_pages(filename: str, data: bytes) -> list[tuple[int, str]]:
    ext = Path(filename).suffix.lower()
    if ext in TEXT_EXT:
        return [(1, data.decode("utf-8", errors="replace"))]
    if ext == ".pdf":
        from pypdf import PdfReader
        reader = PdfReader(io.BytesIO(data))
        pages = [(i + 1, (p.extract_text() or "").strip()) for i, p in enumerate(reader.pages)]
        pages = [(n, t) for n, t in pages if t]
        if not pages:
            raise ValueError("No selectable text found in this PDF (scanned PDF?). "
                             "Upload page images instead so vision OCR can read them.")
        return pages
    if ext == ".docx":
        from docx import Document
        doc = Document(io.BytesIO(data))
        parts = [p.text for p in doc.paragraphs if p.text.strip()]
        for t in doc.tables:
            for r in t.rows:
                parts.append("; ".join(c.text.strip() for c in r.cells if c.text.strip()))
        return [(1, "\n\n".join(parts))]
    if ext == ".csv":
        import pandas as pd
        return [(1, _df_to_text(pd.read_csv(io.BytesIO(data))))]
    if ext == ".xlsx":
        import pandas as pd
        sheets = pd.read_excel(io.BytesIO(data), sheet_name=None)
        return [(i + 1, f"Sheet {name}\n" + _df_to_text(df)) for i, (name, df) in enumerate(sheets.items())]
    if ext in IMAGE_EXT:
        if not llm.available():
            raise ValueError("Image upload needs an LLM for OCR. Configure GEMINI_API_KEY.")
        mime = "image/png" if ext == ".png" else "image/jpeg" if ext in (".jpg", ".jpeg") else "image/webp"
        text = llm.generate("Transcribe all text in this image exactly as written. "
                            "Keep tables as lines of 'column: value'. Output only the transcription.",
                            images=[(data, mime)])
        return [(1, text)]
    raise ValueError(f"Unsupported file type: {ext}. Supported: {', '.join(SUPPORTED)}")
